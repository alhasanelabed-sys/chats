"""Exports preserve report contents without materializing all attendance rows."""
from datetime import date, datetime
import io
import asyncio

import pytest
from openpyxl import load_workbook

from hader import models as m, reports as R
from hader.db import session_scope


@pytest.fixture()
def report_data(client):
    with session_scope() as db:
        emp = m.Employee(emp_code="801", first_name="=2+2", department_id=1)
        db.add(emp)
        db.flush()
        tt = db.query(m.TimeTable).first()
        for day in (date(2026, 8, 2), date(2026, 8, 3)):
            db.add(m.TempSchedule(employee_id=emp.id, att_date=day, timetable_id=tt.id))
        for ts in (datetime(2026, 8, 2, 8, 30), datetime(2026, 8, 2, 16)):
            db.add(m.Transaction(emp_code=emp.emp_code, employee_id=emp.id, punch_time=ts, device_sn=""))
        db.flush()
        yield db


@pytest.mark.parametrize("key", list(R.REPORTS))
def test_stream_and_full_report_match(report_data, key):
    db = report_data
    full = R.build(db, key, date(2026, 8, 2), date(2026, 8, 3), lang="en")
    streamed = R.build(db, key, date(2026, 8, 2), date(2026, 8, 3), lang="en", stream=True)
    assert not isinstance(streamed["rows"], list)
    assert streamed["columns"] == full["columns"]
    assert list(streamed["rows"]) == full["rows"]


def test_streaming_csv_and_excel_preserve_literal_strings(report_data, tmp_path):
    db = report_data
    args = (db, "daily", date(2026, 8, 2), date(2026, 8, 3))
    full = R.build(*args, lang="en")
    streamed = R.build(*args, lang="en", stream=True)
    assert b"".join(R.iter_csv(streamed)) == R.to_csv(full)
    assert b"'=2+2" in R.to_csv(full)
    target = tmp_path / "report.xlsx"
    R.to_xlsx(R.build(*args, lang="en", stream=True), company="=2+2", output=target)
    wb = load_workbook(target, read_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    assert rows[0][0] == "'=2+2"
    assert rows[4][2] == "'=2+2"
    wb.close()


@pytest.mark.parametrize("fmt", ["csv", "xlsx"])
def test_export_endpoint_removes_temporary_file(client, tmp_path, monkeypatch, fmt):
    import tempfile
    monkeypatch.setattr(tempfile, "tempdir", str(tmp_path))
    r = client.get("/api/reports/daily", params={"start": "2026-08-02", "end": "2026-08-03", "fmt": fmt})
    assert r.status_code == 200
    assert r.content.startswith(b"PK" if fmt == "xlsx" else b"\xef\xbb\xbf")
    assert list(tmp_path.glob("hader_export_*")) == []


@pytest.mark.parametrize("failure", ["disconnect", "invalid_range"])
def test_export_file_is_removed_when_delivery_fails(tmp_path, failure):
    from hader.api.system import ExportFileResponse
    path = tmp_path / "export.csv"
    path.write_bytes(b"confidential")
    messages = []

    async def send(message):
        if failure == "disconnect":
            raise OSError("simulated disconnected browser")
        messages.append(message)

    async def receive():
        return {"type": "http.request", "body": b""}

    scope = {"type": "http", "method": "GET", "headers": [], "extensions": {}}
    if failure == "invalid_range":
        scope["headers"] = [(b"range", b"bytes=100-200")]
    response = ExportFileResponse(path)
    if failure == "disconnect":
        with pytest.raises(OSError):
            asyncio.run(response(scope, receive, send))
    else:
        asyncio.run(response(scope, receive, send))
        assert messages[0]["status"] == 416
    assert not path.exists()


@pytest.mark.parametrize("key", ["transactions", "leave"])
@pytest.mark.parametrize("q", ["Sara Ali", "Sales"])
def test_raw_report_search_matches_export(client, key, q):
    with session_scope() as db:
        dep = m.Department(code="S", name="Sales")
        db.add(dep)
        db.flush()
        emp = m.Employee(emp_code="803", first_name="Sara", last_name="Ali", department_id=dep.id)
        db.add(emp)
        db.flush()
        db.add(m.Transaction(emp_code=emp.emp_code, employee_id=emp.id, device_sn="DEVICE_NO_ALIAS",
                             punch_time=datetime(2026, 8, 2, 8)))
        db.add(m.Device(sn="DEVICE_NO_ALIAS", alias=""))
        db.add(m.Leave(employee_id=emp.id, leave_type_id=db.query(m.LeaveType).first().id,
                       start_time=datetime(2026, 8, 2, 8), end_time=datetime(2026, 8, 2, 9)))
    params = {"start": "2026-08-02", "end": "2026-08-02", "q": q, "lang": "en"}
    r = client.get(f"/api/reports/{key}", params=params)
    assert r.status_code == 200
    assert r.json()["total"] == 1 and r.json()["rows"][0]["emp_code"] == "803"
    csv = client.get(f"/api/reports/{key}", params=params | {"fmt": "csv"})
    assert csv.status_code == 200 and b"Sara Ali" in csv.content
    if key == "transactions":
        assert r.json()["rows"][0]["device"] == "DEVICE_NO_ALIAS"
