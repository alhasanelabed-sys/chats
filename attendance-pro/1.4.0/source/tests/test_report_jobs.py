import csv
import io
import time
from datetime import datetime, timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from hader import models as m, store
from hader.api import report_jobs as api
from hader.db import session_scope
from hader.report_jobs import ReportJobManager, validate_options


def wait_job(client, job_id):
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline:
        job = client.get(f'/api/report-jobs/{job_id}').json()
        if job['status'] not in {'queued', 'running'}:
            return job
        time.sleep(.05)
    raise AssertionError('export did not finish within test deadline')


@pytest.mark.parametrize('fmt', ['csv', 'xlsx'])
def test_real_background_export_all_rows_and_no_attendance_mutation(client, tmp_path, monkeypatch, fmt):
    from hader.config import settings
    manager = ReportJobManager(tmp_path / 'jobs', settings.database_url)
    monkeypatch.setattr(api, 'get_manager', lambda: manager)
    employee = client.post('/api/employees', json={'emp_code': 'EXPORT', 'first_name': 'Export'}).json()
    with session_scope() as db:
        for i in range(225):
            db.add(m.Transaction(employee_id=employee['id'], emp_code='EXPORT', device_sn='',
                punch_time=datetime(2026, 8, 2, 8) + timedelta(seconds=i)))
        before_write = store.get(db, 'tcp.write_back')
    try:
        response = client.post('/api/report-jobs', json={'key': 'transactions', 'start': '2026-08-02', 'end': '2026-08-02', 'fmt': fmt})
        assert response.status_code == 202, response.text
        job = wait_job(client, response.json()['id'])
        assert job['status'] == 'passed', job
        assert job['result']['rows'] == 225 and job['downloads']['file']
        downloaded = client.get(f'/api/report-jobs/{job["id"]}/download')
        assert downloaded.status_code == 200 and len(downloaded.content) == job['result']['size_bytes']
        if fmt == 'csv':
            rows = list(csv.reader(io.StringIO(downloaded.content.decode('utf-8-sig'))))
            assert len(rows) == 226 and all(row[0] == 'EXPORT' for row in rows[1:])
        else:
            book = load_workbook(io.BytesIO(downloaded.content), read_only=False)
            assert book.active.max_row == 229
            assert book.active.page_setup.orientation == 'landscape'
            assert book.active.page_setup.fitToWidth == 1 and book.active.print_title_rows == '$1:$4'
            book.close()
        with session_scope() as db:
            assert db.scalar(select(func.count()).select_from(m.Transaction)) == 225
            assert store.get(db, 'tcp.write_back') is before_write
        assert not (manager.root / job['id'] / 'work').exists()
        assert client.delete(f'/api/report-jobs/{job["id"]}').status_code == 200
    finally:
        manager.shutdown()


def test_export_owner_and_expiry_boundary(tmp_path, client, monkeypatch):
    from hader.config import settings
    from hader import report_jobs
    worker = tmp_path / 'blocked.py'
    worker.write_text('import time\ntime.sleep(60)\n')
    manager = ReportJobManager(tmp_path / 'jobs', settings.database_url, worker_path=worker)
    monkeypatch.setattr(api, 'get_manager', lambda: manager)
    monkeypatch.setattr(report_jobs, 'TIMEOUT_SECONDS', .2)
    try:
        job = manager.start({'key': 'daily', 'start': '2026-08-02', 'end': '2026-08-02'}, owner=987)
        with pytest.raises(report_jobs.ReportJobError) as denied:
            manager.get(job['id'], owner=123)
        assert denied.value.status == 404
        result = wait_job(client, job['id'])
        assert result['status'] == 'failed' and '20' in result['error']
        assert not result['downloads']['file'] and not (manager.root / job['id'] / 'work').exists()
    finally:
        manager.shutdown()


def test_background_export_cancel_without_polling_and_permissions(client, tmp_path, monkeypatch):
    from hader.config import settings
    worker = tmp_path / 'blocked.py'; worker.write_text('import time\ntime.sleep(60)\n')
    manager = ReportJobManager(tmp_path / 'jobs', settings.database_url, worker_path=worker)
    monkeypatch.setattr(api, 'get_manager', lambda: manager)
    try:
        r = client.post('/api/report-jobs', json={'start':'2026-08-02','end':'2026-08-02'})
        assert r.status_code == 202
        job = client.post(f'/api/report-jobs/{r.json()["id"]}/cancel').json()
        assert job['status'] == 'cancelled' and not job['downloads']['file']
        client.post('/api/auth/logout')
        assert client.get('/api/report-jobs').status_code == 401
        assert client.get(f'/api/report-jobs/{job["id"]}/download').status_code == 401
    finally:
        manager.shutdown()


@pytest.mark.parametrize('extra', [{'unknown':True}, {'employee_ids':'1,evil'}, {'start':'2026-08-02','end':'2026-08-01'}, {'lang':'x'}, {'fmt':'pdf'}, {'device':'bad\nSN'}])
def test_export_options_cannot_silently_broaden_report(extra):
    with pytest.raises(ValueError):
        validate_options({'start':'2026-08-02','end':'2026-08-02',**extra})
