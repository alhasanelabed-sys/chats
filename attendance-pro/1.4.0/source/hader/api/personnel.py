"""Personnel module: departments, positions, areas, employees, biometrics."""
from __future__ import annotations

import base64
import csv
import io

from fastapi import APIRouter, Body, Depends, File, HTTPException, Request, UploadFile
from fastapi.responses import Response
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session, defer

from ..adms import sync
from ..db import get_db, now
from .. import models as m
from ..photos import photo_bytes, with_photo
from ..terminal.protocol import valid_pin
from .crud import coerce, crud_router
from .deps import audit, page, parse_date, require, ser

router = APIRouter(prefix="/api")


def _dept_dict(db, d):
    out = ser(d)
    out["employees"] = db.scalar(select(func.count()).select_from(m.Employee).where(
        m.Employee.department_id == d.id, m.Employee.status == "active")) or 0
    return out


def _area_dict(db, a):
    out = ser(a)
    out["employees"] = db.scalar(select(func.count()).select_from(m.employee_area).where(
        m.employee_area.c.area_id == a.id)) or 0
    out["devices"] = db.scalar(select(func.count()).select_from(m.Device).where(m.Device.area_id == a.id)) or 0
    return out


router.include_router(crud_router(m.Department, "/departments", "personnel.view", "personnel.edit",
                                  search=("code", "name"), order=m.Department.code, to_dict=_dept_dict))
router.include_router(crud_router(m.Position, "/positions", "personnel.view", "personnel.edit",
                                  search=("code", "name"), order=m.Position.code))
router.include_router(crud_router(m.Area, "/areas", "personnel.view", "personnel.edit",
                                  search=("code", "name"), order=m.Area.code, to_dict=_area_dict))


# --------------------------------------------------------------------------
# Employees
# --------------------------------------------------------------------------

_HIDDEN_EMPLOYEE_FIELDS = {"photo", "portal_hash", "dev_password"}
_PROTECTED_EMPLOYEE_FIELDS = {"portal_hash", "portal_enabled", "portal_must_change", "portal_last_login", "portal_lang"}
_PUBLIC_EMPLOYEE_FIELDS = [c.key for c in m.Employee.__table__.columns if c.key not in _HIDDEN_EMPLOYEE_FIELDS]


def _employee_metrics(db: Session, ids: list[int]) -> tuple[set[int], dict[int, dict[int, int]]]:
    """Read photo presence and biometric counts in batches, never one query per row."""
    photos, counts = set(), {}
    for start in range(0, len(ids), 500):
        batch = ids[start:start + 500]
        photos.update(with_photo(db, batch))
        for emp_id, bio_type, count in db.execute(select(m.BioTemplate.employee_id, m.BioTemplate.bio_type,
                                                         func.count()).where(m.BioTemplate.employee_id.in_(batch))
                                                  .group_by(m.BioTemplate.employee_id, m.BioTemplate.bio_type)):
            counts.setdefault(emp_id, {})[bio_type] = count
    return photos, counts


def emp_dict(db: Session, e: m.Employee, full: bool = False, metrics=None) -> dict:
    # Skip secret/heavy attributes before reading them, including when deferred.
    d = ser(e, fields=_PUBLIC_EMPLOYEE_FIELDS)
    d["name"] = e.display_name
    d["department"] = e.department.label if e.department else ""
    d["position"] = e.position.label if e.position else ""
    d["area_ids"] = [a.id for a in e.areas]
    d["areas"] = ", ".join(a.label for a in e.areas)
    photos, by_employee = metrics if metrics is not None else _employee_metrics(db, [e.id])
    d["has_photo"] = e.id in photos
    counts = by_employee.get(e.id, {})
    d["fp_count"] = counts.get(1, 0)
    d["face_count"] = counts.get(2, 0) + counts.get(9, 0)
    d["palm_count"] = counts.get(8, 0) + counts.get(6, 0)
    d["vein_count"] = counts.get(7, 0)
    return d


def _validate_code(code: str) -> str:
    if not isinstance(code, str):
        raise HTTPException(422, "Employee ID must be text")
    code = code.strip()
    if not valid_pin(code):
        raise HTTPException(422, "Employee ID must be 1-24 Latin letters/digits, starting with a letter or digit")
    return code


@router.get("/employees")
def list_employees(q: str = "", department_id: str = "", area_id: int | None = None, status: str = "active",
                   offset: int = 0, limit: int = 50, db: Session = Depends(get_db),
                   _=Depends(require("personnel.view"))):
    stmt = select(m.Employee).options(defer(m.Employee.photo), defer(m.Employee.portal_hash), defer(m.Employee.dev_password))
    if status in ("active", "resigned"):
        stmt = stmt.where(m.Employee.status == status)
    if q:
        like = f"%{q}%"
        stmt = stmt.where(or_(m.Employee.emp_code.ilike(like), m.Employee.first_name.ilike(like),
                              m.Employee.last_name.ilike(like), m.Employee.card_no.ilike(like),
                              m.Employee.mobile.ilike(like), m.Employee.national_id.ilike(like)))
    if department_id:
        try:
            department_ids = [int(x) for x in department_id.split(",")]
        except ValueError:
            raise HTTPException(422, "invalid department ID list")
        stmt = stmt.where(m.Employee.department_id.in_(department_ids))
    if area_id:
        stmt = stmt.join(m.employee_area).where(m.employee_area.c.area_id == area_id)
    total, rows = page(db, stmt.order_by(m.Employee.emp_code), offset, limit)
    metrics = _employee_metrics(db, [r[0].id for r in rows])
    return {"total": total, "rows": [emp_dict(db, r[0], metrics=metrics) for r in rows]}


@router.get("/employees/{emp_id}")
def get_employee(emp_id: int, db: Session = Depends(get_db), _=Depends(require("personnel.view"))):
    e = db.get(m.Employee, emp_id)
    if not e:
        raise HTTPException(404, "not found")
    d = emp_dict(db, e)
    d["templates"] = [{"id": t.id, "bio_type": t.bio_type, "type": m.BIO_TYPES.get(t.bio_type, str(t.bio_type)),
                       "no": t.bio_no, "index": t.bio_index, "version": f"{t.major_ver}.{t.minor_ver}",
                       "source": t.source_sn, "updated_at": t.updated_at.strftime("%Y-%m-%d %H:%M")}
                      for t in db.scalars(select(m.BioTemplate).where(m.BioTemplate.employee_id == e.id)
                                          .order_by(m.BioTemplate.bio_type, m.BioTemplate.bio_no)).all()]
    d["bio_photos"] = [p.bio_type for p in db.scalars(select(m.BioPhoto).where(m.BioPhoto.employee_id == e.id)).all()]
    return d


def _apply_employee(db: Session, e: m.Employee, data: dict) -> None:
    if _PROTECTED_EMPLOYEE_FIELDS.intersection(data):
        raise HTTPException(422, "Portal credentials must be managed through employee portal actions")
    if "emp_code" in data and data["emp_code"] != e.emp_code:
        raise HTTPException(422, "Employee ID cannot be changed; existing records retain their original ID")
    if "clear_dev_password" in data and type(data["clear_dev_password"]) is not bool:
        raise HTTPException(422, "clear_dev_password must be true or false")
    if "dev_password" in data:
        password = data["dev_password"]
        if not isinstance(password, str) or len(password) > 30 or any(ord(c) < 32 for c in password):
            raise HTTPException(422, "Device password must be up to 30 characters without control characters")
    values = coerce(m.Employee, {k: v for k, v in data.items()
                               if k not in ("photo", "status", "emp_code", "dev_password")})
    if data.get("clear_dev_password") is True:
        values["dev_password"] = ""
    elif data.get("dev_password"):
        values["dev_password"] = data["dev_password"]
    for k, v in values.items():
        setattr(e, k, v)
    if "area_ids" in data:
        ids = [int(x) for x in data.get("area_ids") or []]
        e.areas = list(db.scalars(select(m.Area).where(m.Area.id.in_(ids))).all()) if ids else []


@router.post("/employees")
def create_employee(request: Request, data: dict = Body(...), db: Session = Depends(get_db),
                    _=Depends(require("personnel.edit"))):
    code = _validate_code(data.get("emp_code", ""))
    if db.scalar(select(m.Employee.id).where(m.Employee.emp_code == code)):
        raise HTTPException(409, "Employee ID already exists")
    e = m.Employee(emp_code=code)
    _apply_employee(db, e, data | {"emp_code": code})
    if "area_ids" not in data:
        first_area = db.scalar(select(m.Area).order_by(m.Area.id).limit(1))
        e.areas = [first_area] if first_area else []
    db.add(e)
    db.flush()
    sync.link_transactions(db, e)
    sync.employee_changed(db, e, set(), with_bio=True)
    audit(db, request, "create", "employee", code)
    db.commit()
    return emp_dict(db, e)


@router.put("/employees/{emp_id}")
def update_employee(emp_id: int, request: Request, data: dict = Body(...), db: Session = Depends(get_db),
                    _=Depends(require("personnel.edit"))):
    e = db.get(m.Employee, emp_id)
    if not e:
        raise HTTPException(404, "not found")
    old_areas = {a.id for a in e.areas}
    _apply_employee(db, e, data)
    db.flush()
    sync.employee_changed(db, e, old_areas)
    audit(db, request, "update", "employee", e.emp_code)
    db.commit()
    return emp_dict(db, e)


@router.delete("/employees/{emp_id}")
def delete_employee(emp_id: int, request: Request, db: Session = Depends(get_db),
                    _=Depends(require("personnel.edit"))):
    e = db.get(m.Employee, emp_id)
    if not e:
        raise HTTPException(404, "not found")
    sync.employee_deleted(db, e.emp_code, {a.id for a in e.areas})
    audit(db, request, "delete", "employee", e.emp_code)
    db.delete(e)
    db.commit()
    return {"ok": True}


@router.post("/employees/batch")
def batch_employees(request: Request, data: dict = Body(...), db: Session = Depends(get_db),
                    _=Depends(require("personnel.edit"))):
    """Batch actions: set_department, set_areas, add_areas, resign, reinstate, sync, delete."""
    ids = [int(x) for x in data.get("ids", [])]
    action = data.get("action")
    emps = list(db.scalars(select(m.Employee).where(m.Employee.id.in_(ids))).all())
    n = 0
    for e in emps:
        old_areas = {a.id for a in e.areas}
        if action == "set_department":
            e.department_id = int(data["department_id"])
        elif action == "set_position":
            e.position_id = int(data["position_id"])
        elif action in ("set_areas", "add_areas"):
            new = list(db.scalars(select(m.Area).where(m.Area.id.in_([int(x) for x in data.get("area_ids", [])]))).all())
            e.areas = new if action == "set_areas" else list({a.id: a for a in e.areas + new}.values())
            db.flush()
            n += sync.employee_changed(db, e, old_areas)
            continue
        elif action == "resign":
            e.status = "resigned"
            e.resign_date = parse_date(data.get("resign_date")) or now().date()
            e.resign_type = data.get("resign_type", "")
            e.resign_reason = data.get("resign_reason", "")
            db.flush()
            n += sync.employee_changed(db, e, old_areas)
            continue
        elif action == "reinstate":
            e.status, e.resign_date = "active", None
            db.flush()
            n += sync.employee_changed(db, e, set(), with_bio=True)
            continue
        elif action == "sync":
            for dev in sync.area_devices(db, old_areas):
                n += sync.push_employee_to_device(db, dev, e)
            continue
        elif action == "delete":
            sync.employee_deleted(db, e.emp_code, old_areas)
            db.delete(e)
            continue
        else:
            raise HTTPException(422, "unknown action")
    audit(db, request, "batch." + str(action), "employee", f"{len(emps)} employees")
    db.commit()
    return {"ok": True, "employees": len(emps), "commands": n}


@router.post("/employees/{emp_id}/photo")
async def upload_photo(emp_id: int, request: Request, file: UploadFile = File(...),
                       db: Session = Depends(get_db), _=Depends(require("personnel.edit"))):
    data = await file.read()
    if len(data) > 2_000_000 or not data.startswith(b"\xff\xd8"):
        raise HTTPException(422, "JPEG photo up to 2 MB required")
    e = db.get(m.Employee, emp_id)
    if not e:
        raise HTTPException(404, "not found")
    e.photo = base64.b64encode(data).decode()
    db.flush()
    for dev in sync.area_devices(db, {a.id for a in e.areas}):
        from ..adms import commands as C
        sync.queue(db, dev.sn, C.userpic_update(e.emp_code, e.photo), "Photo " + e.emp_code)
    audit(db, request, "photo", "employee", e.emp_code)
    db.commit()
    return {"ok": True}


@router.get("/employees/{emp_id}/photo")
def get_photo(emp_id: int, db: Session = Depends(get_db), _=Depends(require("personnel.view"))):
    e = db.get(m.Employee, emp_id)
    data = photo_bytes(db, e) if e else None
    if not data:
        raise HTTPException(404, "no photo")
    return Response(data, media_type="image/jpeg", headers={"Cache-Control": "private, max-age=300"})


@router.delete("/employees/{emp_id}/templates/{tpl_id}")
def delete_template(emp_id: int, tpl_id: int, request: Request, db: Session = Depends(get_db),
                    _=Depends(require("personnel.edit"))):
    t = db.get(m.BioTemplate, tpl_id)
    e = db.get(m.Employee, emp_id)
    if not t or not e or t.employee_id != emp_id:
        raise HTTPException(404, "not found")
    for dev in sync.area_devices(db, {a.id for a in e.areas}):
        cmd = sync.template_delete_command(dev, e.emp_code, t.bio_type, t.bio_no, t.bio_index)
        if cmd:
            sync.queue(db, dev.sn, cmd, f"Delete T{t.bio_type} {e.emp_code}")
    db.delete(t)
    audit(db, request, "delete_template", "employee", e.emp_code)
    db.commit()
    return {"ok": True}


# --------------------------------------------------------------------------
# Import / export
# --------------------------------------------------------------------------

IMPORT_FIELDS = ["emp_code", "first_name", "last_name", "department", "position", "card_no", "gender",
                 "hire_date", "mobile", "email", "national_id"]
_ALIASES = {"id": "emp_code", "pin": "emp_code", "code": "emp_code", "رقم": "emp_code", "الرقم": "emp_code",
            "name": "first_name", "الاسم": "first_name", "القسم": "department", "dept": "department",
            "card": "card_no", "البطاقة": "card_no", "الوظيفة": "position"}


def _read_table(filename: str, data: bytes) -> list[dict]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        head = [str(h or "").strip() for h in rows[0]]
        return [{head[i]: ("" if v is None else str(v)) for i, v in enumerate(r) if i < len(head)} for r in rows[1:]]
    text = data.decode("utf-8-sig", errors="replace")
    return list(csv.DictReader(io.StringIO(text)))


@router.post("/employees/import")
async def import_employees(request: Request, file: UploadFile = File(...), db: Session = Depends(get_db),
                           _=Depends(require("personnel.edit"))):
    rows = _read_table(file.filename or "", await file.read())
    depts = {d.name: d for d in db.scalars(select(m.Department)).all()}
    positions = {p.name: p for p in db.scalars(select(m.Position)).all()}
    default_area = db.scalar(select(m.Area).order_by(m.Area.id).limit(1))
    created = updated = 0
    errors = []
    for i, raw in enumerate(rows, start=2):
        row = {}
        for k, v in raw.items():
            key = _ALIASES.get(str(k).strip().lower(), str(k).strip().lower())
            row[key] = (v or "").strip()
        raw_code = row.get("emp_code", "")
        e = db.scalar(select(m.Employee).where(m.Employee.emp_code == raw_code)) if raw_code else None
        try:
            code = e.emp_code if e is not None else _validate_code(raw_code)
        except HTTPException as exc:
            errors.append(f"row {i}: {exc.detail}")
            continue
        if e is None:
            e = db.scalar(select(m.Employee).where(m.Employee.emp_code == code))
        is_new = e is None
        if is_new:
            e = m.Employee(emp_code=code, areas=[default_area] if default_area else [])
            db.add(e)
        dname = row.get("department")
        if dname:
            if dname not in depts:
                depts[dname] = m.Department(code=f"D{len(depts) + 1}-{dname[:20]}", name=dname)
                db.add(depts[dname])
                db.flush()
            e.department = depts[dname]
        elif is_new:
            e.department_id = db.scalar(select(m.Department.id).order_by(m.Department.id).limit(1))
        pname = row.get("position")
        if pname:
            if pname not in positions:
                positions[pname] = m.Position(code=f"P{len(positions) + 1}-{pname[:20]}", name=pname)
                db.add(positions[pname])
                db.flush()
            e.position = positions[pname]
        vals = coerce(m.Employee, {k: row[k] for k in ("first_name", "last_name", "card_no", "gender",
                                                       "hire_date", "mobile", "email", "national_id")
                                   if row.get(k)})
        for k, v in vals.items():
            setattr(e, k, v)
        db.flush()
        if is_new:
            sync.link_transactions(db, e)
        sync.employee_changed(db, e, set() if is_new else {a.id for a in e.areas}, with_bio=is_new)
        created += is_new
        updated += not is_new
    audit(db, request, "import", "employee", f"{created} new, {updated} updated")
    db.commit()
    return {"created": created, "updated": updated, "errors": errors[:50]}


@router.get("/employees-export")
def export_employees(fmt: str = "xlsx", status: str = "active", db: Session = Depends(get_db),
                     _=Depends(require("personnel.view"))):
    from ..reports import to_csv, to_xlsx
    emps = db.scalars(select(m.Employee).options(defer(m.Employee.photo), defer(m.Employee.portal_hash),
                                               defer(m.Employee.dev_password))
                      .where(m.Employee.status == status).order_by(m.Employee.emp_code)).all()
    cols = [("emp_code", "الرقم"), ("first_name", "الاسم الأول"), ("last_name", "اسم العائلة"),
            ("department", "القسم"), ("position", "الوظيفة"), ("card_no", "البطاقة"), ("gender", "الجنس"),
            ("hire_date", "تاريخ التعيين"), ("mobile", "الجوال"), ("email", "البريد"),
            ("national_id", "رقم الهوية"), ("areas", "المناطق")]
    rows = []
    for e in emps:
        rows.append({"emp_code": e.emp_code, "first_name": e.first_name, "last_name": e.last_name,
                     "department": e.department.label if e.department else "",
                     "position": e.position.label if e.position else "", "card_no": e.card_no,
                     "gender": e.gender, "hire_date": e.hire_date.isoformat() if e.hire_date else "",
                     "mobile": e.mobile, "email": e.email, "national_id": e.national_id,
                     "areas": ", ".join(a.label for a in e.areas)})
    rep = {"title": "Employees", "columns": [{"key": k, "label": k if fmt == "csv" else lbl} for k, lbl in cols],
           "rows": rows}
    if fmt == "csv":
        return Response(to_csv(rep), media_type="text/csv",
                        headers={"Content-Disposition": "attachment; filename=employees.csv"})
    return Response(to_xlsx(rep), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": "attachment; filename=employees.xlsx"})


@router.get("/lookups")
def lookups(db: Session = Depends(get_db), _=Depends(require("personnel.view"))):
    """Small reference lists for the UI's select boxes."""
    return {
        "departments": [{"id": d.id, "name": d.label, "code": d.code, "parent_id": d.parent_id}
                        for d in db.scalars(select(m.Department).order_by(m.Department.code)).all()],
        "positions": [{"id": p.id, "name": p.label} for p in db.scalars(select(m.Position).order_by(m.Position.code)).all()],
        "areas": [{"id": a.id, "name": a.label} for a in db.scalars(select(m.Area).order_by(m.Area.code)).all()],
        "timetables": [{"id": t.id, "name": t.alias, "color": t.color, "check_in": t.check_in.strftime("%H:%M"),
                        "check_out": t.check_out.strftime("%H:%M")}
                       for t in db.scalars(select(m.TimeTable).order_by(m.TimeTable.alias)).all()],
        "shifts": [{"id": s.id, "name": s.alias} for s in db.scalars(select(m.Shift).order_by(m.Shift.alias)).all()],
        "leave_types": [{"id": t.id, "name": t.label, "code": t.code, "color": t.color}
                        for t in db.scalars(select(m.LeaveType).order_by(m.LeaveType.code)).all()],
        "devices": [{"sn": d.sn, "name": d.label} for d in db.scalars(select(m.Device).order_by(m.Device.alias)).all()],
        "roles": [{"id": r.id, "name": r.name} for r in db.scalars(select(m.Role).order_by(m.Role.name)).all()],
    }


@router.get("/employees-search")
def search_employees(q: str = "", limit: int = 20, db: Session = Depends(get_db),
                     _=Depends(require("personnel.view"))):
    like = f"%{q}%"
    rows = db.scalars(select(m.Employee).where(or_(
        m.Employee.emp_code.ilike(like), m.Employee.first_name.ilike(like), m.Employee.last_name.ilike(like)))
        .order_by(m.Employee.emp_code).limit(limit)).all()
    return [{"id": e.id, "emp_code": e.emp_code, "name": e.display_name,
             "department": e.department.label if e.department else ""} for e in rows]
