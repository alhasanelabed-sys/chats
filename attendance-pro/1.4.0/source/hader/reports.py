"""Report definitions and CSV / Excel export."""
from __future__ import annotations

import csv
import io
from itertools import islice
from datetime import date, datetime, time, timedelta

from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session, defer, lazyload

from . import models as m
from .adms.protocol import VERIFY_TYPES
from .engine import Engine, employee_day_page, employee_query, iter_attendance, summarize

STATUS_AR = {
    "present": "حاضر", "late": "متأخر", "early": "خروج مبكر", "late_early": "متأخر وخروج مبكر",
    "absent": "غائب", "partial_absent": "غياب جزئي", "leave": "إجازة", "holiday": "عطلة رسمية", "off": "راحة",
    "incomplete": "بصمة ناقصة", "unscheduled": "بدون جدول", "pending": "لم يحضر بعد", "": "",
}
STATUS_EN = {
    "present": "Present", "late": "Late", "early": "Early leave", "late_early": "Late & early",
    "absent": "Absent", "partial_absent": "Partial absence", "leave": "Leave", "holiday": "Holiday", "off": "Day off",
    "incomplete": "Missed punch", "unscheduled": "Unscheduled", "pending": "Not yet", "": "",
}
STATE_AR = {0: "دخول", 1: "خروج", 2: "خروج استراحة", 3: "عودة استراحة", 4: "دخول إضافي", 5: "خروج إضافي", 255: "-"}
STATE_EN = {0: "Check-In", 1: "Check-Out", 2: "Break-Out", 3: "Break-In", 4: "OT-In", 5: "OT-Out", 255: "-"}
VERIFY_AR = {"password": "كلمة مرور", "fingerprint": "بصمة إصبع", "card": "بطاقة", "face": "وجه",
             "palm": "كف", "finger_vein": "وريد الإصبع", "other": "أخرى"}
MONTH_SYMBOL = {"present": "✓", "late": "L", "early": "E", "late_early": "LE", "absent": "A", "partial_absent": "PA",
                "leave": "V", "holiday": "H", "off": "-", "incomplete": "!", "unscheduled": "•", "pending": "", "": ""}

# key -> (Arabic title, English title, kind)
REPORTS = {
    "transactions": ("سجل الحركات", "Transactions", "punch"),
    "first_last": ("أول وآخر بصمة", "First & Last", "day"),
    "time_card": ("بطاقة الدوام", "Time Card", "day"),
    "daily": ("الحضور اليومي", "Daily Attendance", "day"),
    "late": ("تقرير التأخير", "Late Report", "day"),
    "early": ("تقرير الخروج المبكر", "Early Leave", "day"),
    "absent": ("تقرير الغياب", "Absence Report", "day"),
    "overtime": ("تقرير العمل الإضافي", "Overtime Report", "day"),
    "exception": ("البصمات الناقصة", "Missed Punch", "day"),
    "leave": ("تقرير الإجازات", "Leave Report", "leave"),
    "summary": ("الملخص الشهري", "Monthly Summary", "summary"),
    "monthly_status": ("كشف الحضور الشهري", "Monthly Status", "matrix"),
    "department": ("ملخص الأقسام", "Department Summary", "dept"),
}


def hm(mins) -> str:
    try:
        mins = int(mins or 0)
    except (TypeError, ValueError):
        return str(mins)
    if not mins:
        return ""
    return f"{mins // 60}:{mins % 60:02d}"


def _t(d: str) -> str:
    return d[11:16] if d else ""


def _cols(lang: str, spec: list[tuple[str, str, str]]) -> list[dict]:
    return [{"key": k, "label": ar if lang == "ar" else en} for k, ar, en in spec]


def _raw_employee_search(q: str, *, punch: bool = False):
    pattern = f"%{q.strip()}%"
    conditions = [m.Employee.emp_code.ilike(pattern), m.Employee.first_name.ilike(pattern),
                  m.Employee.last_name.ilike(pattern), m.Employee.name_en.ilike(pattern),
                  m.Employee.card_no.ilike(pattern),
                  (m.Employee.first_name + " " + m.Employee.last_name).ilike(pattern),
                  m.Employee.department.has(or_(m.Department.name.ilike(pattern),
                                                m.Department.name_en.ilike(pattern)))]
    if punch:
        conditions.append(m.Transaction.emp_code.ilike(pattern))
    return or_(*conditions)


def _build_full(db: Session, key: str, start: date, end: date, *, lang: str = "ar",
                employee_ids=None, department_ids=None, device_sn: str | None = None,
                stream: bool = False, search: str = "") -> dict:
    """Returns {"title", "columns": [{key,label}], "rows": [dict], "kind"}."""
    if key not in REPORTS:
        raise KeyError(key)
    title_ar, title_en, kind = REPORTS[key]
    status_names = STATUS_AR if lang == "ar" else STATUS_EN
    base = [("emp_code", "الرقم", "ID"), ("name", "الاسم", "Name"), ("department", "القسم", "Department")]
    out = {"title": title_ar if lang == "ar" else title_en, "kind": kind, "key": key,
           "start": start.isoformat(), "end": end.isoformat()}

    if kind == "punch":
        q = select(m.Transaction, m.Employee).outerjoin(m.Employee, m.Employee.id == m.Transaction.employee_id).where(
            m.Transaction.punch_time >= datetime.combine(start, time.min),
            m.Transaction.punch_time < datetime.combine(end + timedelta(days=1), time.min))
        if employee_ids is not None:
            q = q.where(m.Transaction.employee_id.in_(employee_ids))
        if department_ids is not None:
            q = q.where(m.Employee.department_id.in_(department_ids))
        if device_sn:
            q = q.where(m.Transaction.device_sn == device_sn)
        if search.strip():
            q = q.where(_raw_employee_search(search, punch=True))
        aliases = dict(db.execute(select(m.Device.sn, m.Device.alias)).all())
        q = q.order_by(m.Transaction.punch_time, m.Transaction.id).options(
            defer(m.Employee.photo), defer(m.Employee.portal_hash), defer(m.Employee.dev_password),
            lazyload(m.Employee.areas), lazyload(m.Employee.position)).execution_options(yield_per=1000)

        def rows():
            for t, e in db.execute(q):
                v = VERIFY_TYPES.get(t.verify_type, "other")
                yield {
                    "emp_code": t.emp_code, "name": e.display_name if e else "",
                    "department": e.department.label if e and e.department else "",
                    "date": t.punch_time.strftime("%Y-%m-%d"), "time": t.punch_time.strftime("%H:%M:%S"),
                    "state": (STATE_AR if lang == "ar" else STATE_EN).get(t.punch_state, str(t.punch_state)),
                    "verify": VERIFY_AR.get(v, v) if lang == "ar" else v.replace("_", " ").title(),
                    "device": aliases.get(t.device_sn) or t.device_sn or ("يدوي" if lang == "ar" else "Manual"),
                    "temperature": t.temperature if t.temperature is not None else "",
                }
        out["columns"] = _cols(lang, base + [("date", "التاريخ", "Date"), ("time", "الوقت", "Time"),
                                             ("state", "الحالة", "State"), ("verify", "طريقة التحقق", "Verify"),
                                             ("device", "الجهاز", "Device"), ("temperature", "الحرارة", "Temp.")])
        out["rows"] = rows() if stream else list(rows())
        return out

    if kind == "leave":
        q = select(m.Leave, m.Employee, m.LeaveType).join(m.Employee, m.Employee.id == m.Leave.employee_id).join(
            m.LeaveType, m.LeaveType.id == m.Leave.leave_type_id).where(
            m.Leave.start_time < datetime.combine(end + timedelta(days=1), time.min),
            m.Leave.end_time > datetime.combine(start, time.min))
        if employee_ids is not None:
            q = q.where(m.Leave.employee_id.in_(employee_ids))
        if department_ids is not None:
            q = q.where(m.Employee.department_id.in_(department_ids))
        if search.strip():
            q = q.where(_raw_employee_search(search))
        st = {"approved": "معتمدة", "pending": "بانتظار الموافقة", "rejected": "مرفوضة"} if lang == "ar" else {}
        q = q.order_by(m.Leave.start_time, m.Leave.id).options(
            defer(m.Employee.photo), defer(m.Employee.portal_hash), defer(m.Employee.dev_password),
            lazyload(m.Employee.areas), lazyload(m.Employee.position)).execution_options(yield_per=1000)
        rows = ({
            "emp_code": e.emp_code, "name": e.display_name, "department": e.department.label if e.department else "",
            "type": lt.name, "start": lv.start_time.strftime("%Y-%m-%d %H:%M"),
            "end": lv.end_time.strftime("%Y-%m-%d %H:%M"),
            "days": round((lv.end_time - lv.start_time).total_seconds() / 86400, 2),
            "status": st.get(lv.status, lv.status), "reason": lv.reason,
        } for lv, e, lt in db.execute(q))
        out["columns"] = _cols(lang, base + [("type", "النوع", "Type"), ("start", "من", "From"),
                                             ("end", "إلى", "To"), ("days", "أيام", "Days"),
                                             ("status", "الحالة", "Status"), ("reason", "السبب", "Reason")])
        out["rows"] = rows if stream else list(rows)
        return out

    days = iter_attendance(db, start, end, employee_ids=employee_ids, department_ids=department_ids, q=search)

    if kind == "summary":
        rows = summarize(days)
        for r in rows:
            r["leave_detail"] = ", ".join(f"{k}:{v}" for k, v in sorted(r.pop("leave_by_code").items()))
            for k in ("required", "worked", "late", "early", "absent_minutes", "ot", "leave"):
                r[k + "_hm"] = hm(r[k])
        out["columns"] = _cols(lang, base + [
            ("scheduled_days", "أيام الدوام", "Work days"), ("present_days", "أيام الحضور", "Present"),
            ("absent_days", "أيام الغياب", "Absent"), ("partial_absent_days", "أيام الغياب الجزئي", "Partial absence days"),
            ("absent_minutes", "دقائق الغياب", "Absent minutes"), ("late_count", "مرات التأخير", "Late times"),
            ("late_hm", "مدة التأخير", "Late"), ("early_count", "مرات الخروج المبكر", "Early times"),
            ("early_hm", "مدة الخروج المبكر", "Early"), ("missed_punch", "بصمات ناقصة", "Missed"),
            ("required_hm", "الساعات المطلوبة", "Required"), ("worked_hm", "ساعات العمل", "Worked"),
            ("ot_hm", "العمل الإضافي", "Overtime"), ("leave_days", "أيام الإجازة", "Leave days"),
            ("leave_detail", "تفصيل الإجازات", "Leave detail"), ("holidays", "العطل", "Holidays"),
            ("off_days", "أيام الراحة", "Days off")])
        out["rows"] = rows
        return out

    if kind == "matrix":
        dates = []
        d = start
        while d <= end:
            dates.append(d)
            d += timedelta(days=1)
        by_emp: dict[int, dict] = {}
        for r in days:
            row = by_emp.setdefault(r.employee_id, {"emp_code": r.emp_code, "name": r.name,
                                                     "department": r.department, "_p": 0, "_a": 0, "_l": 0})
            sym = MONTH_SYMBOL.get(r.status, "")
            if r.status == "leave" and r.leave_codes:
                sym = r.leave_codes[0]
            row[r.att_date.isoformat()] = sym
            row["_p"] += 1 if r.status in ("present", "late", "early", "late_early", "incomplete") else 0
            row["_a"] += 1 if r.status in ("absent", "partial_absent") else 0
            row["_l"] += 1 if r.late else 0
        cols = base + [(d.isoformat(), f"{d.day}", f"{d.day}") for d in dates]
        cols += [("_p", "حضور", "P"), ("_a", "غياب", "A"), ("_l", "تأخير", "L")]
        out["columns"] = _cols(lang, cols)
        out["rows"] = list(by_emp.values())
        out["legend"] = {v: (STATUS_AR if lang == "ar" else STATUS_EN)[k] for k, v in MONTH_SYMBOL.items() if v}
        return out

    if kind == "dept":
        per: dict[str, dict] = {}
        for s in summarize(days):
            d = per.setdefault(s["department"] or "-", {"department": s["department"] or "-", "employees": 0,
                                                         "present_days": 0.0, "absent_days": 0, "partial_absent_days": 0,
                                                         "absent_minutes": 0, "late_count": 0,
                                                         "late": 0, "early": 0, "worked": 0, "ot": 0, "leave_days": 0})
            d["employees"] += 1
            for k in ("present_days", "absent_days", "partial_absent_days", "absent_minutes", "late_count", "late", "early", "worked", "ot", "leave_days"):
                d[k] += s[k]
        rows = []
        for d in per.values():
            d["present_days"] = round(d["present_days"], 2)
            for k in ("late", "early", "worked", "ot"):
                d[k + "_hm"] = hm(d[k])
            rows.append(d)
        out["columns"] = _cols(lang, [
            ("department", "القسم", "Department"), ("employees", "الموظفون", "Employees"),
            ("present_days", "أيام الحضور", "Present days"), ("absent_days", "أيام الغياب", "Absent days"),
            ("partial_absent_days", "غياب جزئي", "Partial absence days"), ("absent_minutes", "دقائق الغياب", "Absent minutes"),
            ("late_count", "مرات التأخير", "Late times"), ("late_hm", "مدة التأخير", "Late"),
            ("early_hm", "الخروج المبكر", "Early"), ("worked_hm", "ساعات العمل", "Worked"),
            ("ot_hm", "العمل الإضافي", "Overtime"), ("leave_days", "أيام الإجازة", "Leave days")])
        out["rows"] = rows
        return out

    # ---- day based reports
    filt = {
        "late": lambda r: r.late > 0,
        "early": lambda r: r.early > 0,
        "absent": lambda r: r.status in ("absent", "partial_absent"),
        "overtime": lambda r: r.ot > 0,
        "exception": lambda r: any(x in ("missed_in", "missed_out") for x in r.exceptions),
        "first_last": lambda r: bool(r.punches),
        "time_card": lambda r: True,
        "daily": lambda r: True,
    }[key]
    def rows():
        for r in days:
            if not filt(r):
                continue
            d = r.as_dict()
            d["clock_in"], d["clock_out"] = _t(d["clock_in"]), _t(d["clock_out"])
            d["sched"] = f"{_t(d['sched_in'])}-{_t(d['sched_out'])}" if d["sched_in"] else ""
            d["status_label"] = status_names.get(r.status, r.status)
            if r.holiday:
                d["status_label"] += f" ({r.holiday})"
            if r.leave_codes:
                d["status_label"] += " [" + ",".join(r.leave_codes) + "]"
            d["punch_list"] = " ".join(d["punches"])
            d["first"] = d["punches"][0] if d["punches"] else ""
            d["last"] = d["punches"][-1] if len(d["punches"]) > 1 else ""
            for k in ("required", "worked", "late", "early", "absent", "ot", "leave"):
                d[k + "_hm"] = hm(d[k])
            missed = {"missed_in": "بدون دخول" if lang == "ar" else "No check-in",
                      "missed_out": "بدون خروج" if lang == "ar" else "No check-out"}
            d["missed"] = "، ".join(missed[x] for x in r.exceptions if x in missed)
            yield d
    day_cols = [("date", "التاريخ", "Date")] + base
    spec = {
        "first_last": day_cols + [("first", "أول بصمة", "First punch"), ("last", "آخر بصمة", "Last punch"),
                                  ("punch_list", "كل البصمات", "All punches")],
        "time_card": day_cols + [("timetable", "الدوام", "Timetable"), ("sched", "الوقت المقرر", "Schedule"),
                                 ("punch_list", "البصمات", "Punches"), ("worked_hm", "ساعات العمل", "Worked"),
                                 ("status_label", "الحالة", "Status")],
        "daily": day_cols + [("timetable", "الدوام", "Timetable"), ("sched", "الوقت المقرر", "Schedule"),
                             ("clock_in", "الدخول", "Clock in"), ("clock_out", "الخروج", "Clock out"),
                             ("late_hm", "تأخير", "Late"), ("early_hm", "خروج مبكر", "Early"),
                             ("absent_hm", "غياب", "Absent"), ("worked_hm", "عمل", "Worked"),
                             ("ot_hm", "إضافي", "OT"), ("leave_hm", "إجازة", "Leave"),
                             ("status_label", "الحالة", "Status")],
        "late": day_cols + [("sched", "الوقت المقرر", "Schedule"), ("clock_in", "الدخول", "Clock in"),
                            ("late_hm", "مدة التأخير", "Late")],
        "early": day_cols + [("sched", "الوقت المقرر", "Schedule"), ("clock_out", "الخروج", "Clock out"),
                             ("early_hm", "مدة الخروج المبكر", "Early")],
        "absent": day_cols + [("timetable", "الدوام", "Timetable"), ("sched", "الوقت المقرر", "Schedule"),
                              ("absent_hm", "مدة الغياب", "Absent")],
        "overtime": day_cols + [("sched", "الوقت المقرر", "Schedule"), ("clock_in", "الدخول", "Clock in"),
                                ("clock_out", "الخروج", "Clock out"), ("ot_hm", "العمل الإضافي", "Overtime"),
                                ("status_label", "الحالة", "Status")],
        "exception": day_cols + [("sched", "الوقت المقرر", "Schedule"), ("clock_in", "الدخول", "Clock in"),
                                 ("clock_out", "الخروج", "Clock out"), ("missed", "النقص", "Missing")],
    }[key]
    out["columns"] = _cols(lang, spec)
    out["rows"] = rows() if stream else list(rows())
    return out


def _matches(row: dict, q: str, status: str) -> bool:
    if q:
        hay = " ".join(str(row.get(k, "")) for k in ("emp_code", "name", "department")).casefold()
        if q.casefold() not in hay:
            return False
    return not status or row.get("status") == status


def _day_page(db: Session, key: str, start: date, end: date, *, lang: str,
              employee_ids=None, department_ids=None, offset: int = 0, limit: int = 200,
              q: str = "", status: str = "") -> dict:
    """Stream day rows through the engine and retain only the requested page."""
    if key not in ("first_last", "time_card", "daily", "late", "early", "absent", "overtime", "exception"):
        raise KeyError(key)
    title_ar, title_en, kind = REPORTS[key]
    names = STATUS_AR if lang == "ar" else STATUS_EN
    base = [("emp_code", "الرقم", "ID"), ("name", "الاسم", "Name"), ("department", "القسم", "Department")]
    filt = {"late": lambda r: r.late > 0, "early": lambda r: r.early > 0,
            "absent": lambda r: r.status in ("absent", "partial_absent"),
            "overtime": lambda r: r.ot > 0,
            "exception": lambda r: any(x in ("missed_in", "missed_out") for x in r.exceptions),
            "first_last": lambda r: bool(r.punches), "time_card": lambda r: True, "daily": lambda r: True}[key]
    rows: list[dict] = []
    total = 0
    direct_page = key in ("daily", "time_card") and not status
    filters = dict(employee_ids=employee_ids, department_ids=department_ids, q=q)
    if direct_page:
        total, days = employee_day_page(db, start, end, offset=offset, limit=limit, **filters)
    else:
        days = iter_attendance(db, start, end, **filters)
    for r in days:
        if not filt(r) or status and r.status != status:
            continue
        if (direct_page or total >= offset) and len(rows) < limit:
            d = r.as_dict()
            d["clock_in"], d["clock_out"] = _t(d["clock_in"]), _t(d["clock_out"])
            d["sched"] = f"{_t(d['sched_in'])}-{_t(d['sched_out'])}" if d["sched_in"] else ""
            d["status_label"] = names.get(r.status, r.status)
            if r.holiday: d["status_label"] += f" ({r.holiday})"
            if r.leave_codes: d["status_label"] += " [" + ",".join(r.leave_codes) + "]"
            d["punch_list"] = " ".join(d["punches"])
            d["first"] = d["punches"][0] if d["punches"] else ""
            d["last"] = d["punches"][-1] if len(d["punches"]) > 1 else ""
            for k in ("required", "worked", "late", "early", "absent", "ot", "leave"):
                d[k + "_hm"] = hm(d[k])
            missed = {"missed_in": "بدون دخول" if lang == "ar" else "No check-in",
                      "missed_out": "بدون خروج" if lang == "ar" else "No check-out"}
            d["missed"] = "، ".join(missed[x] for x in r.exceptions if x in missed)
            rows.append(d)
        if not direct_page:
            total += 1
    day_cols = [("date", "التاريخ", "Date")] + base
    spec = {
        "first_last": day_cols + [("first", "أول بصمة", "First punch"), ("last", "آخر بصمة", "Last punch"), ("punch_list", "كل البصمات", "All punches")],
        "time_card": day_cols + [("timetable", "الدوام", "Timetable"), ("sched", "الوقت المقرر", "Schedule"), ("punch_list", "البصمات", "Punches"), ("worked_hm", "ساعات العمل", "Worked"), ("status_label", "الحالة", "Status")],
        "daily": day_cols + [("timetable", "الدوام", "Timetable"), ("sched", "الوقت المقرر", "Schedule"), ("clock_in", "الدخول", "Clock in"), ("clock_out", "الخروج", "Clock out"), ("late_hm", "تأخير", "Late"), ("early_hm", "خروج مبكر", "Early"), ("absent_hm", "غياب", "Absent"), ("worked_hm", "عمل", "Worked"), ("ot_hm", "إضافي", "OT"), ("leave_hm", "إجازة", "Leave"), ("status_label", "الحالة", "Status")],
        "late": day_cols + [("sched", "الوقت المقرر", "Schedule"), ("clock_in", "الدخول", "Clock in"), ("late_hm", "مدة التأخير", "Late")],
        "early": day_cols + [("sched", "الوقت المقرر", "Schedule"), ("clock_out", "الخروج", "Clock out"), ("early_hm", "مدة الخروج المبكر", "Early")],
        "absent": day_cols + [("timetable", "الدوام", "Timetable"), ("sched", "الوقت المقرر", "Schedule"), ("absent_hm", "مدة الغياب", "Absent")],
        "overtime": day_cols + [("sched", "الوقت المقرر", "Schedule"), ("clock_in", "الدخول", "Clock in"), ("clock_out", "الخروج", "Clock out"), ("ot_hm", "العمل الإضافي", "Overtime"), ("status_label", "الحالة", "Status")],
        "exception": day_cols + [("sched", "الوقت المقرر", "Schedule"), ("clock_in", "الدخول", "Clock in"), ("clock_out", "الخروج", "Clock out"), ("missed", "النقص", "Missing")],
    }[key]
    return {"title": title_ar if lang == "ar" else title_en, "kind": kind, "key": key,
            "start": start.isoformat(), "end": end.isoformat(), "columns": _cols(lang, spec),
            "rows": rows, "total": total, "offset": offset, "limit": limit,
            "has_more": offset + len(rows) < total}


def _punch_or_leave_page(db: Session, key: str, start: date, end: date, *, lang: str,
                         employee_ids=None, department_ids=None, device_sn=None,
                         offset: int = 0, limit: int = 200, q: str = "") -> dict:
    off, lim = max(0, offset), max(1, min(limit, 1000))
    base = [("emp_code", "الرقم", "ID"), ("name", "الاسم", "Name"), ("department", "القسم", "Department")]
    if key == "transactions":
        stmt = select(m.Transaction, m.Employee).outerjoin(m.Employee, m.Employee.id == m.Transaction.employee_id).where(
            m.Transaction.punch_time >= datetime.combine(start, time.min),
            m.Transaction.punch_time < datetime.combine(end + timedelta(days=1), time.min)).options(
                defer(m.Employee.photo), defer(m.Employee.portal_hash), defer(m.Employee.dev_password),
                lazyload(m.Employee.areas), lazyload(m.Employee.position))
        if employee_ids is not None: stmt = stmt.where(m.Transaction.employee_id.in_(employee_ids))
        if department_ids is not None: stmt = stmt.where(m.Employee.department_id.in_(department_ids))
        if device_sn: stmt = stmt.where(m.Transaction.device_sn == device_sn)
        if q:
            stmt = stmt.where(_raw_employee_search(q, punch=True))
        count = db.scalar(select(func.count()).select_from(stmt.order_by(None).subquery())) or 0
        aliases = dict(db.execute(select(m.Device.sn, m.Device.alias)).all())
        rows = []
        for t, e in db.execute(stmt.order_by(m.Transaction.punch_time, m.Transaction.id).offset(off).limit(lim)).all():
            v = VERIFY_TYPES.get(t.verify_type, "other")
            rows.append({"emp_code": t.emp_code, "name": e.display_name if e else "", "department": e.department.label if e and e.department else "",
                         "date": t.punch_time.strftime("%Y-%m-%d"), "time": t.punch_time.strftime("%H:%M:%S"),
                         "state": (STATE_AR if lang == "ar" else STATE_EN).get(t.punch_state, str(t.punch_state)),
                         "verify": VERIFY_AR.get(v, v) if lang == "ar" else v.replace("_", " ").title(),
                         "device": aliases.get(t.device_sn) or t.device_sn or ("يدوي" if lang == "ar" else "Manual"),
                         "temperature": t.temperature if t.temperature is not None else ""})
        spec = base + [("date", "التاريخ", "Date"), ("time", "الوقت", "Time"), ("state", "الحالة", "State"),
                       ("verify", "طريقة التحقق", "Verify"), ("device", "الجهاز", "Device"), ("temperature", "الحرارة", "Temp.")]
    else:
        stmt = select(m.Leave, m.Employee, m.LeaveType).join(m.Employee, m.Employee.id == m.Leave.employee_id).join(
            m.LeaveType, m.LeaveType.id == m.Leave.leave_type_id).where(
                m.Leave.start_time < datetime.combine(end + timedelta(days=1), time.min),
                m.Leave.end_time > datetime.combine(start, time.min)).options(
                    defer(m.Employee.photo), defer(m.Employee.portal_hash), defer(m.Employee.dev_password),
                    lazyload(m.Employee.areas), lazyload(m.Employee.position))
        if employee_ids is not None: stmt = stmt.where(m.Leave.employee_id.in_(employee_ids))
        if department_ids is not None: stmt = stmt.where(m.Employee.department_id.in_(department_ids))
        if q:
            stmt = stmt.where(_raw_employee_search(q))
        count = db.scalar(select(func.count()).select_from(stmt.order_by(None).subquery())) or 0
        st = {"approved": "معتمدة", "pending": "بانتظار الموافقة", "rejected": "مرفوضة"} if lang == "ar" else {}
        rows = [{"emp_code": e.emp_code, "name": e.display_name, "department": e.department.label if e.department else "",
                 "type": lt.name, "start": lv.start_time.strftime("%Y-%m-%d %H:%M"), "end": lv.end_time.strftime("%Y-%m-%d %H:%M"),
                 "days": round((lv.end_time - lv.start_time).total_seconds() / 86400, 2), "status": st.get(lv.status, lv.status), "reason": lv.reason}
                for lv, e, lt in db.execute(stmt.order_by(m.Leave.start_time, m.Leave.id).offset(off).limit(lim)).all()]
        spec = base + [("type", "النوع", "Type"), ("start", "من", "From"), ("end", "إلى", "To"),
                       ("days", "أيام", "Days"), ("status", "الحالة", "Status"), ("reason", "السبب", "Reason")]
    title_ar, title_en, kind = REPORTS[key]
    return {"title": title_ar if lang == "ar" else title_en, "kind": kind, "key": key,
            "start": start.isoformat(), "end": end.isoformat(), "columns": _cols(lang, spec), "rows": rows,
            "total": count, "offset": off, "limit": lim, "has_more": off + len(rows) < count}


def _employee_ids_page(db: Session, start: date, end: date, employee_ids, department_ids,
                       q: str, offset: int, limit: int):
    stmt = employee_query(start, end, employee_ids=employee_ids, department_ids=department_ids,
                          q=q).with_only_columns(m.Employee.id)
    total = db.scalar(select(func.count()).select_from(stmt.order_by(None).subquery())) or 0
    ids = list(db.scalars(stmt.offset(max(0, offset)).limit(max(1, min(limit, 1000)))).all())
    return total, ids


def _matrix_page(db: Session, start: date, end: date, *, lang: str, employee_ids=None,
                 department_ids=None, offset: int = 0, limit: int = 200, q: str = "") -> dict:
    total, ids = _employee_ids_page(db, start, end, employee_ids, department_ids, q, offset, limit)
    days = Engine(db, start, end, employee_ids=ids, department_ids=department_ids).run()
    by_emp: dict[int, dict] = {}
    for r in days:
        row = by_emp.setdefault(r.employee_id, {"emp_code": r.emp_code, "name": r.name,
                                                 "department": r.department, "_p": 0, "_a": 0, "_l": 0})
        row[r.att_date.isoformat()] = MONTH_SYMBOL.get(r.status, "")
        if r.status == "leave" and r.leave_codes:
            row[r.att_date.isoformat()] = r.leave_codes[0]
        row["_p"] += 1 if r.status in ("present", "late", "early", "late_early", "incomplete") else 0
        row["_a"] += 1 if r.status in ("absent", "partial_absent") else 0
        row["_l"] += 1 if r.late else 0
    dates = []
    d = start
    while d <= end:
        dates.append(d); d += timedelta(days=1)
    base = [("emp_code", "الرقم", "ID"), ("name", "الاسم", "Name"), ("department", "القسم", "Department")]
    cols = base + [(d.isoformat(), str(d.day), str(d.day)) for d in dates] + [("_p", "حضور", "P"), ("_a", "غياب", "A"), ("_l", "تأخير", "L")]
    title_ar, title_en, kind = REPORTS["monthly_status"]
    return {"title": title_ar if lang == "ar" else title_en, "kind": kind, "key": "monthly_status",
            "start": start.isoformat(), "end": end.isoformat(), "columns": _cols(lang, cols),
            "rows": list(by_emp.values()), "total": total, "offset": max(0, offset), "limit": limit,
            "has_more": max(0, offset) + len(ids) < total,
            "legend": {v: (STATUS_AR if lang == "ar" else STATUS_EN)[k] for k, v in MONTH_SYMBOL.items() if v}}


def _summary_page(db: Session, start: date, end: date, *, lang: str, employee_ids=None,
                  department_ids=None, offset: int = 0, limit: int = 200, q: str = "") -> dict:
    total, ids = _employee_ids_page(db, start, end, employee_ids, department_ids, q, offset, limit)
    rows = summarize(Engine(db, start, end, employee_ids=ids, department_ids=department_ids).iter_rows())
    for r in rows:
        r["leave_detail"] = ", ".join(f"{k}:{v}" for k, v in sorted(r.pop("leave_by_code").items()))
        for k in ("required", "worked", "late", "early", "absent_minutes", "ot", "leave"):
            r[k + "_hm"] = hm(r[k])
    base = [("emp_code", "الرقم", "ID"), ("name", "الاسم", "Name"), ("department", "القسم", "Department")]
    spec = base + [("scheduled_days", "أيام الدوام", "Work days"), ("present_days", "أيام الحضور", "Present"),
                   ("absent_days", "أيام الغياب", "Absent"), ("partial_absent_days", "أيام الغياب الجزئي", "Partial absence days"),
                   ("absent_minutes", "دقائق الغياب", "Absent minutes"), ("late_count", "مرات التأخير", "Late times"),
                   ("late_hm", "مدة التأخير", "Late"), ("early_count", "مرات الخروج المبكر", "Early times"),
                   ("early_hm", "مدة الخروج المبكر", "Early"), ("missed_punch", "بصمات ناقصة", "Missed"),
                   ("required_hm", "الساعات المطلوبة", "Required"), ("worked_hm", "ساعات العمل", "Worked"),
                   ("ot_hm", "العمل الإضافي", "Overtime"), ("leave_days", "أيام الإجازة", "Leave days"),
                   ("leave_detail", "تفصيل الإجازات", "Leave detail"), ("holidays", "العطل", "Holidays"),
                   ("off_days", "أيام الراحة", "Days off")]
    title_ar, title_en, kind = REPORTS["summary"]
    return {"title": title_ar if lang == "ar" else title_en, "kind": kind, "key": "summary",
            "start": start.isoformat(), "end": end.isoformat(), "columns": _cols(lang, spec), "rows": rows,
            "total": total, "offset": max(0, offset), "limit": limit,
            "has_more": max(0, offset) + len(ids) < total}


def _build_stream(db: Session, key: str, start: date, end: date, *, lang: str,
                  employee_ids=None, department_ids=None, device_sn=None, q: str = "") -> dict:
    kind = REPORTS[key][2]
    kw = dict(lang=lang, employee_ids=employee_ids, department_ids=department_ids,
              device_sn=device_sn, stream=True)
    if kind in ("punch", "leave", "dept"):
        out = _build_full(db, key, start, end, **kw, search=q)
        out["rows"] = iter(out["rows"])
        return out
    # Matrix and summary calculations are bounded to a small employee batch.
    ids = list(db.scalars(employee_query(start, end, employee_ids=employee_ids,
                                         department_ids=department_ids, q=q).with_only_columns(m.Employee.id)))
    out = _build_full(db, key, start, end, **(kw | {"employee_ids": []}))

    def rows():
        for i in range(0, len(ids), 50):
            batch = _build_full(db, key, start, end, **(kw | {"employee_ids": ids[i:i + 50]}))
            yield from batch["rows"]

    out["rows"] = rows()
    return out


def build(db: Session, key: str, start: date, end: date, *, lang: str = "ar",
          employee_ids=None, department_ids=None, device_sn: str | None = None,
          offset: int = 0, limit: int | None = None, q: str = "", status: str = "",
          stream: bool = False) -> dict:
    """Build a report page with a truthful total.

    The report definitions remain compatible with callers that request all rows
    (``limit=None``), while API calls use a bounded page by default.
    """
    if stream:
        if limit is not None:
            raise ValueError("streaming exports require limit=None")
        out = _build_stream(db, key, start, end, lang=lang, employee_ids=employee_ids,
                            department_ids=department_ids, device_sn=device_sn, q=q[:100])
        out["rows"] = islice((r for r in out["rows"] if _matches(r, "", status)), max(0, offset), None)
        return out
    if limit is not None and key in ("first_last", "time_card", "daily", "late", "early", "absent", "overtime", "exception"):
        return _day_page(db, key, start, end, lang=lang, employee_ids=employee_ids,
                         department_ids=department_ids, offset=max(0, int(offset or 0)),
                         limit=max(1, min(int(limit or 200), 1000)), q=q[:100], status=status[:40])
    if limit is not None and key in ("transactions", "leave"):
        return _punch_or_leave_page(db, key, start, end, lang=lang, employee_ids=employee_ids,
                                    department_ids=department_ids, device_sn=device_sn,
                                    offset=max(0, int(offset or 0)), limit=max(1, min(int(limit or 200), 1000)), q=q[:100])
    if limit is not None and key == "monthly_status":
        return _matrix_page(db, start, end, lang=lang, employee_ids=employee_ids,
                            department_ids=department_ids, offset=max(0, int(offset or 0)),
                            limit=max(1, min(int(limit or 200), 1000)), q=q[:100])
    if limit is not None and key == "summary":
        return _summary_page(db, start, end, lang=lang, employee_ids=employee_ids,
                             department_ids=department_ids, offset=max(0, int(offset or 0)),
                             limit=max(1, min(int(limit or 200), 1000)), q=q[:100])
    out = _build_full(db, key, start, end, lang=lang, employee_ids=employee_ids,
                      department_ids=department_ids, device_sn=device_sn, search=q[:100])
    if status:
        out["rows"] = [r for r in out.get("rows", []) if _matches(r, "", status)]
    rows = out.get("rows", [])
    total = len(rows)
    off = max(0, int(offset or 0))
    if limit is None:
        page_rows = rows[off:]
        lim = len(page_rows)
    else:
        lim = max(1, min(int(limit or 200), 1000))
        page_rows = rows[off:off + lim]
    out["rows"] = page_rows
    out.update({"total": total, "offset": off, "limit": lim, "has_more": off + len(page_rows) < total})
    return out


# --------------------------------------------------------------------------
# Export
# --------------------------------------------------------------------------

def iter_csv(report: dict):
    def safe(v):
        if isinstance(v, (list, dict)):
            v = str(v)
        if isinstance(v, str) and v.lstrip().startswith(("=", "+", "-", "@")):
            return "'" + v
        return v
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c["label"] for c in report["columns"]])
    yield ("\ufeff" + buf.getvalue()).encode("utf-8")
    buf.seek(0)
    buf.truncate(0)
    for r in report["rows"]:
        w.writerow([safe(r.get(c["key"], "")) for c in report["columns"]])
        yield buf.getvalue().encode("utf-8")
        buf.seek(0)
        buf.truncate(0)


def to_csv(report: dict) -> bytes:
    return b"".join(iter_csv(report))


def to_xlsx(report: dict, company: str = "", rtl: bool = True, *, output=None) -> bytes | None:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    def safe(v):
        if isinstance(v, (list, dict)):
            v = str(v)
        if isinstance(v, str) and v.lstrip().startswith(("=", "+", "-", "@")):
            return "'" + v
        return v

    # Write-only worksheets keep memory bounded and allow a report to be split at
    # Excel's 1,048,576-row limit instead of producing a corrupt workbook.
    wb = Workbook(write_only=True)
    head_fill = PatternFill("solid", fgColor="1F6FB2")
    columns = report["columns"]
    sheet_no, row_no = 0, 0
    ws = None

    def new_sheet():
        nonlocal sheet_no, row_no
        sheet_no += 1; row_no = 0
        sheet = wb.create_sheet(title=(report["title"][:25] + (f"_{sheet_no}" if sheet_no > 1 else ""))[:31])
        sheet.sheet_view.rightToLeft = rtl
        sheet.page_setup.orientation = "landscape" if len(columns) > 7 else "portrait"
        sheet.page_setup.paperSize = "9"  # ISO A4; write-only sheets expose no paper constants
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = "1:4"
        for i, c in enumerate(columns, 1):
            sheet.column_dimensions[get_column_letter(i)].width = min(45, max(8, len(str(c["label"])) + 2))
        sheet.append([safe(company)])
        sheet.append([f"{report['title']}   {report.get('start', '')} → {report.get('end', '')}"])
        blank = [None] * max(0, len(columns) - 1)
        sheet.append(blank)
        header = []
        for c in columns:
            cell = WriteOnlyCell(sheet, value=c["label"])
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            header.append(cell)
        sheet.append(header)
        return sheet

    ws = new_sheet()
    for r in report["rows"]:
        if row_no >= 1_048_572:
            ws.print_area = f"A1:{get_column_letter(len(columns))}{row_no + 4}"
            ws = new_sheet()
        values = [safe(r.get(c["key"], "")) for c in columns]
        ws.append(values)
        row_no += 1
    ws.print_area = f"A1:{get_column_letter(len(columns))}{row_no + 4}"
    if output is not None:
        wb.save(output)
        return None
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
