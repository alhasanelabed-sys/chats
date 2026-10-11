import io
import pytest
from datetime import date, datetime, timedelta
from html.parser import HTMLParser

from hader import models as m
from hader.db import session_scope
from hader.printing import write_document


class PrintTable(HTMLParser):
    def __init__(self, document):
        super().__init__()
        self.headers, self.rows, self.widths = [], [], []
        self.table_count = 0
        self.row = self.cell = None
        self.feed(document)

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == 'table':
            self.table_count += 1
        if tag == 'th' and 'data-key' in attrs:
            self.headers.append(attrs['data-key'])
        if tag == 'col':
            self.widths.append((attrs['data-key'], float(attrs['style'].removeprefix('width:').removesuffix('%'))))
        if tag == 'tr' and 'data-row' in attrs:
            self.row = []
        if tag == 'td' and self.row is not None:
            self.cell = [attrs['data-key'], '']

    def handle_data(self, data):
        if self.cell is not None:
            self.cell[1] += data

    def handle_endtag(self, tag):
        if tag == 'td' and self.cell is not None:
            self.row.append(self.cell)
            self.cell = None
        if tag == 'tr' and self.row is not None:
            self.rows.append(self.row)
            self.row = None


def test_print_includes_every_row_beyond_screen_page_and_filters(client):
    employee = client.post('/api/employees', json={'emp_code': 'PRINT', 'first_name': '<script>alert(1)</script>'}).json()
    with session_scope() as db:
        for i in range(225):
            db.add(m.Transaction(employee_id=employee['id'], emp_code='PRINT', device_sn='',
                punch_time=datetime(2026, 8, 2, 8) + timedelta(seconds=i), source='import'))
        db.add(m.Transaction(employee_id=employee['id'], emp_code='PRINT', device_sn='', punch_time=datetime(2026, 8, 3, 8)))
    screen = client.get('/api/reports/transactions?start=2026-08-02&end=2026-08-02').json()
    assert screen['total'] == 225 and len(screen['rows']) == 200
    printed = client.get('/api/reports/transactions/print?start=2026-08-02&end=2026-08-02&orientation=portrait')
    assert printed.status_code == 200 and printed.headers['x-print-rows'] == '225'
    assert 'data-row="225"' in printed.text and 'data-row="226"' not in printed.text
    assert 'A4 portrait' in printed.text and '<script>alert(1)</script>' not in printed.text
    assert '&lt;script&gt;alert(1)&lt;/script&gt;' in printed.text
    assert printed.headers['cache-control'] == 'no-store'
    assert printed.headers['x-print-layout'] == 'fit-width'
    assert printed.headers['x-print-paper'] == 'A4'


def test_wide_month_print_preserves_every_column_and_identity():
    columns = [{'key': k, 'label': k} for k in ['emp_code', 'name', 'department'] + [f'day_{i}' for i in range(1, 32)]]
    out = io.StringIO()
    metadata = write_document({'title': 'Month', 'columns': columns, 'rows': [{col['key']: col['key'] for col in columns}]}, out)
    assert metadata['column_bands'] == 1
    assert metadata['sections'] == 1
    for column in columns:
        assert f'data-key="{column["key"]}"' in out.getvalue()
    assert out.getvalue().count('<th data-key="emp_code">') == 1
    assert out.getvalue().count('<th data-key="name">') == 1
    assert 'overflow:hidden' not in out.getvalue() and 'table-layout:fixed' in out.getvalue()


@pytest.mark.parametrize('days', [30, 31])
@pytest.mark.parametrize('orientation,paper', [('auto', 'A4'), ('portrait', 'A4'), ('landscape', 'A3')])
def test_month_all_columns_share_each_row_and_stream_last_employee(days, orientation, paper):
    dates = [(date(2026, 8, 1) + timedelta(days=i)).isoformat() for i in range(days)]
    keys = ['emp_code', 'name', 'department'] + dates + ['_p', '_a', '_l']
    columns = [{'key': key, 'label': key} for key in keys]
    rows = ({key: f'{key}:employee-{i}' for key in keys} for i in range(1, 204))
    out = io.StringIO()
    metadata = write_document({'title': 'Month', 'columns': columns, 'rows': rows}, out,
                              orientation=orientation, paper=paper)
    document = out.getvalue()
    parsed = PrintTable(document)
    assert metadata['rows'] == 203 and metadata['column_bands'] == 1
    assert parsed.table_count == 1 and parsed.headers == keys
    assert len(parsed.rows) == 203
    assert all([key for key, value in row] == keys for row in parsed.rows)
    assert parsed.rows[-1][-1] == ['_l', '_l:employee-203']
    assert len(parsed.widths) == len(keys)
    assert sum(width for _, width in parsed.widths) == pytest.approx(100, abs=.001)
    assert dict(parsed.widths)['name'] > dict(parsed.widths)[dates[0]] > 0
    assert 'display:table-header-group' in document
    assert 'data-layout="fit-width"' in document
    assert f'size:{paper} {"landscape" if orientation == "auto" else orientation}' in document


def test_print_legend_and_all_text_are_escaped():
    out = io.StringIO()
    write_document({'title': '<script>title()</script>', 'start': '<img src=x>', 'end': '2026-08-31',
                    'columns': [{'key': 'name', 'label': '<svg onload=header()>'}],
                    'rows': [{'name': '<img onerror=name()>'}], 'legend': {'<X>': '<b>leave</b>'}},
                   out, company='<script>company()</script>')
    document = out.getvalue()
    assert '<script>title()' not in document and '<script>company()' not in document
    assert '<svg onload=header()>' not in document and '<img onerror=name()>' not in document
    assert '&lt;X&gt;: &lt;b&gt;leave&lt;/b&gt;' in document


def test_month_excel_preserves_all_columns_and_one_page_width():
    from hader.reports import to_xlsx
    from openpyxl import load_workbook
    dates = [(date(2026, 8, 1) + timedelta(days=i)).isoformat() for i in range(31)]
    keys = ['emp_code', 'name', 'department'] + dates + ['_p', '_a', '_l']
    columns = [{'key': key, 'label': key} for key in keys]
    report = {'title': 'Month', 'kind': 'matrix', 'columns': columns,
              'rows': [{key: f'{key}:last' for key in keys}]}
    book = load_workbook(io.BytesIO(to_xlsx(report)))
    try:
        sheet = book.active
        assert sheet.page_setup.orientation == 'landscape'
        assert sheet.page_setup.fitToWidth == 1 and sheet.page_setup.fitToHeight == 0
        assert sheet.sheet_properties.pageSetUpPr.fitToPage is True
        assert sheet.print_title_rows == '$1:$4'
        assert sheet.max_column == len(keys)
        assert list(next(sheet.iter_rows(min_row=5, values_only=True))) == [f'{key}:last' for key in keys]
        assert sheet.column_dimensions['B'].width > sheet.column_dimensions['D'].width
        assert sheet.column_dimensions['D'].width < 8
    finally:
        book.close()


def test_real_month_print_uses_all_filtered_employees_beyond_screen_limit(client):
    with session_scope() as db:
        db.add_all([m.Employee(emp_code=f'PRINTMONTH{i:03}', first_name=f'Employee {i}')
                    for i in range(1, 204)])
        db.add(m.Employee(emp_code='EXCLUDED', first_name='Other employee'))
    params = {'start': '2026-08-01', 'end': '2026-08-31', 'q': 'PRINTMONTH'}
    screen = client.get('/api/reports/monthly_status', params=params).json()
    assert screen['total'] == 203 and len(screen['rows']) == 200
    printed = client.get('/api/reports/monthly_status/print', params=params)
    assert printed.status_code == 200, printed.text
    assert printed.headers['x-print-rows'] == '203'
    assert printed.headers['x-print-orientation'] == 'landscape'
    parsed = PrintTable(printed.text)
    assert parsed.table_count == 1 and len(parsed.headers) == 37
    assert parsed.headers[3] == '2026-08-01' and parsed.headers[-4] == '2026-08-31'
    assert len(parsed.rows) == 203 and all(len(row) == 37 for row in parsed.rows)
    assert parsed.rows[-1][0] == ['emp_code', 'PRINTMONTH203']
    assert 'EXCLUDED' not in printed.text


def test_print_direction_validation_and_permissions(client):
    assert client.get('/api/reports/daily/print?orientation=diagonal').status_code == 422
    assert client.get('/api/reports/daily/print?paper=letter').status_code == 422
    paper = client.get('/api/reports/daily/print?start=2026-08-02&end=2026-08-02&paper=A3&orientation=landscape')
    assert paper.status_code == 200 and paper.headers['x-print-paper'] == 'A3'
    assert 'A3 landscape' in paper.text
    assert client.get('/api/reports/missing/print').status_code == 404
    client.post('/api/auth/logout')
    assert client.get('/api/reports/daily/print').status_code == 401
